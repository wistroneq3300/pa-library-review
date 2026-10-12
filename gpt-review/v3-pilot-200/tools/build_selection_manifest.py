#!/usr/bin/env python3
"""Build the fixed 200-case sampling manifest without making V3 judgments."""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
V2_PATH = ROOT / "gpt-review" / "test_gpt_review.json"
XLSX_PATH = ROOT / "data" / "REVISED_commands_merged_with_raw.xlsx"
OUTPUT_PATH = ROOT / "gpt-review" / "v3-pilot-200" / "SELECTION_MANIFEST.json"

ALIASES = {
    "REQUIRES PACKAGE": "REQUIRES PACKAGE / USER CONFIRMATION",
    "USER CONFIRMATION": "REQUIRES PACKAGE / USER CONFIRMATION",
}

TARGETS = {
    "Functionality": {
        "FULLY AUTOMATABLE": 12,
        "REQUIRES PACKAGE / USER CONFIRMATION": 32,
        "MANUAL ONLY": 16,
        "BLOCKED": 20,
    },
    "Reliability": {
        "FULLY AUTOMATABLE": 2,
        "REQUIRES PACKAGE / USER CONFIRMATION": 15,
        "MANUAL ONLY": 7,
        "BLOCKED": 6,
    },
    "Performance": {"REQUIRES PACKAGE / USER CONFIRMATION": 25},
    "Compatibility": {
        "FULLY AUTOMATABLE": 4,
        "REQUIRES PACKAGE / USER CONFIRMATION": 18,
        "MANUAL ONLY": 4,
        "BLOCKED": 4,
    },
    "Stability": {
        "REQUIRES PACKAGE / USER CONFIRMATION": 6,
        "MANUAL ONLY": 12,
        "BLOCKED": 7,
    },
    "(No Main Function)": {
        "REQUIRES PACKAGE / USER CONFIRMATION": 6,
        "BLOCKED": 4,
    },
}

CALIBRATION = [
    ("Functionality", 29, "Wistron-HW-00032-V004", "Locked CRPS vendor-spec versus BLOCKED boundary"),
    ("Functionality", 67, "Wistron-HW-00070-V004", "Locked PSU firmware package/approval boundary"),
    ("Functionality", 150, "Wistron-HW-00153-V002", "Locked hybrid OS-install/TPM boundary and source conflict"),
    ("Functionality", 198, "Wistron-HW-00207-V002", "Locked manual serial-communication boundary"),
    ("Functionality", 201, "Wistron-HW-00210-V002", "Locked hybrid USB BIOS/OS coverage boundary"),
    ("Functionality", 79, "Wistron-HW-00082-V004", "CPU stress duration and worker-count preservation"),
    ("Functionality", 87, "Wistron-HW-00090-V004", "Memory stress command semantics and eight-hour duration"),
    ("Functionality", 2260, "Wistron-BMC-00848-V003", "Pure read-only BMC firmware inventory"),
    ("Reliability", 2, "Wistron-System RAS-00374-V003", "Read-only ACPI BERT/ERST/HEST evidence"),
    ("Reliability", 4, "Wistron-System RAS-00378-V003", "High-risk DIMM error-injection classification"),
    ("Performance", 3, "Wistron-Performance CPU-00002-V003", "Vendor HPL package and real workload intent"),
    ("Performance", 4, "Wistron-Performance DIMM-00003-V004", "STREAM memory-bandwidth workload intent"),
    ("Performance", 15, "Wistron-Performance Networking-00030-V004", "InfiniBand peer and fabric performance scope"),
    ("Compatibility", 76, "Wistron-Storage-00059-V004", "Destructive FIO target and false-completion risk"),
    ("Compatibility", 251, "Wistron-NV GPU-00244-V002", "NVIDIA PCIe Gen5 enumeration and link scope"),
    ("Compatibility", 257, "Wistron-NV GPU-00250-V002", "NVLink/NVSwitch topology versus PCIe distinction"),
    ("Stability", 68, "Wistron-Long Term Stress-00067-V004", "Manual AC-cycle 1000 boundary"),
    ("Stability", 9, "Wistron-Margin-00008-V002", "GPU burn under controlled thermal/voltage margin"),
    ("(No Main Function)", 3, "Wistron---00356-V002", "Candidate genuine LTSSM semantic blocker"),
    ("(No Main Function)", 7, "Wistron---00842-V003", "BMC firmware downgrade package/recovery boundary"),
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_ref(value: Any) -> tuple[str, int]:
    if isinstance(value, dict):
        return str(value["sheet"]), int(value["source_row"])
    match = re.match(r"@\{sheet=(.*?);\s*source_row=(\d+)(?:;.*?)?\}", str(value))
    if not match:
        raise ValueError(f"Unsupported source_ref: {value!r}")
    return match.group(1), int(match.group(2))


def evenly_spaced(rows: list[dict[str, Any]], count: int) -> list[dict[str, Any]]:
    if count == 0:
        return []
    if len(rows) < count:
        raise ValueError(f"Need {count} rows but only {len(rows)} are available")
    if count == 1:
        return [rows[len(rows) // 2]]
    indices = [round(i * (len(rows) - 1) / (count - 1)) for i in range(count)]
    if len(set(indices)) != count:
        raise AssertionError("Systematic sampling generated duplicate indices")
    return [rows[index] for index in indices]


def main() -> None:
    v2 = json.loads(V2_PATH.read_text(encoding="utf-8"))["cases"]
    by_key: dict[tuple[str, int, str], dict[str, Any]] = {}
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for case in v2:
        sheet, row = parse_ref(case["source_ref"])
        key = (sheet, row, case["code"])
        if key in by_key:
            raise ValueError(f"Duplicate V2 key {key}")
        historical = str(case["automation_classification"])
        normalized = ALIASES.get(historical, historical)
        record = {
            "sheet": sheet,
            "source_row": row,
            "code": case["code"],
            "test_name": case["test_name"],
            "historical_v2_classification": historical,
            "sampling_classification": normalized,
            "historical_v2_risk_level": case.get("risk_level"),
            "historical_v2_destructive_actions": case.get("destructive_actions"),
        }
        by_key[key] = record
        grouped[(sheet, normalized)].append(record)

    workbook = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=False)
    xlsx_keys = set()
    for sheet_name in workbook.sheetnames:
        rows = workbook[sheet_name].iter_rows(values_only=True)
        next(rows)
        for physical_row, values in enumerate(rows, start=2):
            xlsx_keys.add((sheet_name, physical_row, str(values[0] or "").strip()))

    selected: list[dict[str, Any]] = []
    selected_keys: set[tuple[str, int, str]] = set()
    for sheet, row, code, reason in CALIBRATION:
        key = (sheet, row, code)
        if key not in by_key or key not in xlsx_keys:
            raise KeyError(f"Calibration key does not resolve exactly: {key}")
        item = dict(by_key[key])
        item.update({"calibration": True, "selection_reason": reason})
        selected.append(item)
        selected_keys.add(key)

    current = Counter((item["sheet"], item["sampling_classification"]) for item in selected)
    for sheet, class_targets in TARGETS.items():
        for classification, target in class_targets.items():
            need = target - current[(sheet, classification)]
            candidates = [
                item
                for item in sorted(
                    grouped[(sheet, classification)], key=lambda value: value["source_row"]
                )
                if (item["sheet"], item["source_row"], item["code"]) not in selected_keys
            ]
            for item in evenly_spaced(candidates, need):
                copy = dict(item)
                copy.update(
                    {
                        "calibration": False,
                        "selection_reason": (
                            "Systematic source-order sample for sheet/classification coverage; "
                            "historical labels are sampling strata, not accepted V3 conclusions"
                        ),
                    }
                )
                selected.append(copy)
                selected_keys.add((copy["sheet"], copy["source_row"], copy["code"]))

    calibration = selected[: len(CALIBRATION)]
    expansion = sorted(
        selected[len(CALIBRATION) :],
        key=lambda value: (list(TARGETS).index(value["sheet"]), value["source_row"], value["code"]),
    )
    selected = calibration + expansion
    for sequence, item in enumerate(selected, start=1):
        canonical = f'{item["sheet"]}|{item["source_row"]}|{item["code"]}'
        item["sequence"] = sequence
        item["case_variant_id"] = hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:16]
        item["source_key"] = canonical

    if len(selected) != 200 or len(selected_keys) != 200:
        raise AssertionError(f"Expected 200 unique cases, got {len(selected)}/{len(selected_keys)}")

    actual = Counter((item["sheet"], item["sampling_classification"]) for item in selected)
    expected = Counter(
        (sheet, classification)
        for sheet, targets in TARGETS.items()
        for classification, count in targets.items()
        for _ in range(count)
    )
    if actual != expected:
        raise AssertionError(f"Target distribution mismatch: {actual} != {expected}")

    output = {
        "schema_version": "v3-selection-manifest-1.0",
        "generated_at": "2026-10-13T00:00:00+08:00",
        "baseline_commit": "770c2bd81d329414f308999faa55d20212210fab",
        "source_sha256": {
            "xlsx": sha256(XLSX_PATH),
            "gpt_v2": sha256(V2_PATH),
        },
        "selection_method": (
            "Twenty manually chosen calibration boundary/domain cases followed by deterministic "
            "even-spacing within fixed sheet and historical-V2-classification strata. Historical "
            "classification and risk fields are used only to ensure coverage and BLOCKED/high-risk "
            "weighting; they are not accepted as V3 judgments. The pilot is not a representative "
            "estimate of all 3,112 cases."
        ),
        "target_distribution": TARGETS,
        "calibration_count": len(CALIBRATION),
        "total_count": len(selected),
        "cases": selected,
    }
    OUTPUT_PATH.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "output": str(OUTPUT_PATH.relative_to(ROOT)),
                "count": len(selected),
                "unique": len(selected_keys),
                "by_sheet": Counter(item["sheet"] for item in selected),
                "by_class": Counter(item["sampling_classification"] for item in selected),
                "sha256": sha256(OUTPUT_PATH),
            },
            ensure_ascii=False,
            indent=2,
            default=dict,
        )
    )


if __name__ == "__main__":
    main()

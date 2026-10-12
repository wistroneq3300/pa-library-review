#!/usr/bin/env python3
"""Read one testcase directly from the source XLSX and cross-reference JSON reviews.

This utility only locates and prints source material. It does not classify,
revise, or write testcase review decisions.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
XLSX = ROOT / "data" / "REVISED_commands_merged_with_raw.xlsx"
DEEPSEEK = ROOT / "data" / "tests.json"
GPT_V2 = ROOT / "gpt-review" / "test_gpt_review.json"


def source_ref(value: Any) -> tuple[str | None, int | None]:
    if isinstance(value, dict):
        return value.get("sheet"), int(value.get("source_row"))
    if isinstance(value, str):
        text = value.strip().removeprefix("@{").removesuffix("}")
        fields = dict(
            part.strip().split("=", 1) for part in text.split(";") if "=" in part
        )
        row = fields.get("source_row")
        return fields.get("sheet"), int(row) if row else None
    return None, None


def load_deepseek() -> dict[tuple[str, int, str], dict[str, Any]]:
    data = json.loads(DEEPSEEK.read_text(encoding="utf-8"))
    index: dict[tuple[str, int, str], dict[str, Any]] = {}
    for sheet in data["sheets"].values():
        name = sheet["name"]
        for physical_row, item in enumerate(sheet["items"], start=2):
            index[(name, physical_row, str(item.get("code", "")))] = item
    return index


def load_v2() -> list[dict[str, Any]]:
    data = json.loads(GPT_V2.read_text(encoding="utf-8"))
    return data["cases"]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--code", required=True)
    parser.add_argument("--sheet")
    parser.add_argument("--row", type=int, help="Physical XLSX row, including header row 1")
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=False)
    matches: list[dict[str, Any]] = []
    for name in workbook.sheetnames:
        if args.sheet and name != args.sheet:
            continue
        sheet = workbook[name]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        for physical_row, values in enumerate(rows, start=2):
            if args.row and physical_row != args.row:
                continue
            if str(values[0] or "").strip() != args.code:
                continue
            matches.append(
                {
                    "source_key": {
                        "sheet": name,
                        "source_row": physical_row,
                        "code": args.code,
                    },
                    "xlsx": dict(zip(headers, values, strict=True)),
                    "original_row_sha256": hashlib.sha256(
                        json.dumps(
                            dict(zip(headers, values, strict=True)),
                            ensure_ascii=False,
                            sort_keys=True,
                            separators=(",", ":"),
                            default=str,
                        ).encode("utf-8")
                    ).hexdigest(),
                }
            )

    deepseek = load_deepseek()
    v2_cases = load_v2()
    for match in matches:
        key = match["source_key"]
        tuple_key = (key["sheet"], key["source_row"], key["code"])
        match["tests_json_row"] = deepseek.get(tuple_key)
        v2_matches = []
        for case in v2_cases:
            if case.get("code") != args.code:
                continue
            sheet_name, row_number = source_ref(case.get("source_ref"))
            if sheet_name == key["sheet"] and row_number == key["source_row"]:
                v2_matches.append(case)
        match["gpt_v2_review"] = v2_matches

    print(json.dumps({"query": vars(args), "matches": matches}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

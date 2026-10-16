#!/usr/bin/env python3
"""Read-only V3.1 full-review integrity and ED-REVIEW-001 quality gates."""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
WORKSPACE = ROOT / "gpt-review" / "v3-1-full-3112"
MANIFEST_PATH = WORKSPACE / "SELECTION_MANIFEST.json"
PROGRESS_PATH = WORKSPACE / "PROGRESS.json"
RECORDS_PATH = WORKSPACE / "REVIEW_RECORDS"
AUDIT_RECORDS_PATH = ROOT / "gpt-review" / "v3-1-final-audit" / "AUDIT_RECORDS"
RETROSPECTIVE_PATH = WORKSPACE / "RETROSPECTIVE_IMPACT_20261015.md"
DECISIONS_PATH = WORKSPACE / "ENGINEERING_DECISIONS_APPROVED_20261015.md"

PROTECTED_HASHES = {
    ROOT / "data" / "REVISED_commands_merged_with_raw.xlsx": "e2e737beb9cc4141f89a353813a48ab8034f55da56facba86a72fb9501a114d8",
    ROOT / "data" / "tests.json": "5eefd664be9ed794e4429095f9d8e0c969d392a688b787af71f4f319526fdd37",
    ROOT / "data" / "tests_gpt_merged.json": "34d0133ce0bab7647179f9822174838ecdfd5ee90e766eb9802eb8672dcfe4a2",
}
XLSX = ROOT / "data" / "REVISED_commands_merged_with_raw.xlsx"
PUBLISHED_891_CHECKPOINT = "b8da580a50d36ed1538d1c5262c3c95bb5bbcab2"

CLASSES = {
    "FULLY AUTOMATABLE",
    "REQUIRES PACKAGE / USER CONFIRMATION",
    "MANUAL ONLY",
    "BLOCKED",
}
OUTCOMES = {
    "KEEP", "IMPROVED", "RECLASSIFIED", "SOURCE_CONFLICT",
    "NEED_ENGINEER_DECISION",
}
SECOND_REVIEW_OUTCOMES = {"AGREE", "IMPROVED", "CHANGED", "BLOCKED", "UNRESOLVED"}
REQUIRED_FIELDS = {
    "schema_version", "review_sequence", "library_sequence", "source_key",
    "case_variant_id", "source_ref", "source_integrity", "source_snapshot",
    "historical_reviews", "original_test_intent", "original_method",
    "automation_classification", "execution_readiness", "execution_status",
    "tool_result", "engineer_verdict", "manual_activity", "required_inputs",
    "safety_requirements", "agent_execution_scope", "engineer_actions",
    "command_applicability", "modernized_recommendation",
    "project_override_considerations", "expected_evidence",
    "source_conflict_analysis", "blocked_analysis", "review_improvement",
    "independent_technical_reasoning", "source_provenance", "no_automatic_retry",
    "review_outcome", "second_review_outcome", "regression_result",
    "per_case_quality_gate",
}
ANTI_TEMPLATE_FIELDS = (
    "original_test_intent", "manual_activity", "review_improvement",
    "independent_technical_reasoning",
)
LEGACY_IDENTICAL_DUPLICATE_FILES = {
    "0518_39402f59a448e3b6.json", "0519_6b34ec093c6e1b2b.json",
    "0520_c9451c9f6eb7ea64.json", "0521_2d51ebe47f6901d6.json",
    "0525_b6c29c9e289d4341.json", "0552_6654944d20fd967e.json",
    "0563_d74ef06918f01855.json", "0564_dcac9dc29a5de253.json",
    "0567_65981bef4012edee.json", "0573_f7add07f26553b08.json",
    "0575_8076bf5583797f2e.json", "0577_8159b8a39771b7cc.json",
    "0580_bd0fcccf4ed2ed3a.json", "0604_a66a28a910b311d1.json",
    "0605_1ad807356a7b6821.json", "0607_05b60839d481825b.json",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_json_strict(path: Path, failures: list[str], warnings: list[str]) -> Any:
    duplicates: list[tuple[str, Any, Any]] = []

    def object_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        result: dict[str, Any] = {}
        for key, value in pairs:
            if key in result:
                duplicates.append((key, result[key], value))
            result[key] = value
        return result

    try:
        value = json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=object_pairs)
    except (OSError, json.JSONDecodeError) as exc:
        failures.append(f"{path.relative_to(ROOT)}: invalid JSON: {exc}")
        return None
    for key, first, second in duplicates:
        allowed = (
            path.parent == RECORDS_PATH
            and path.name in LEGACY_IDENTICAL_DUPLICATE_FILES
            and key == "engineering_decision_update"
            and first == second
        )
        if allowed:
            warnings.append(f"legacy identical duplicate key retained: {path.name}:{key}")
        else:
            failures.append(f"{path.relative_to(ROOT)}: duplicate JSON key {key!r}")
    return value


def main() -> int:
    failures: list[str] = []
    warnings: list[str] = []
    for path, expected in PROTECTED_HASHES.items():
        if not path.exists() or sha256(path) != expected:
            failures.append(f"protected source hash changed: {path.relative_to(ROOT)}")

    manifest = load_json_strict(MANIFEST_PATH, failures, warnings) or {"cases": []}
    progress = load_json_strict(PROGRESS_PATH, failures, warnings) or {}
    cases = manifest.get("cases", [])
    if len(cases) != 3112:
        failures.append(f"manifest count is {len(cases)}, expected 3112")
    by_key = {case.get("source_key"): case for case in cases}
    if len(by_key) != len(cases):
        failures.append("manifest source keys are not unique")
    if len({case.get("case_variant_id") for case in cases}) != len(cases):
        failures.append("manifest case variants are not unique")

    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=False)
    direct_keys: list[str] = []
    try:
        for sheet_name in workbook.sheetnames:
            rows = workbook[sheet_name].iter_rows(values_only=True)
            next(rows)
            for physical_row, values in enumerate(rows, start=2):
                direct_keys.append(f"{sheet_name}|{physical_row}|{str(values[0] or '').strip()}")
    finally:
        workbook.close()
    if direct_keys != [case.get("source_key") for case in cases]:
        failures.append("manifest order or source identity differs from direct XLSX order")

    audit_records: list[dict[str, Any]] = []
    for path in sorted(AUDIT_RECORDS_PATH.glob("*.json")):
        payload = load_json_strict(path, failures, warnings)
        if isinstance(payload, list):
            audit_records.extend(item for item in payload if isinstance(item, dict))
        else:
            failures.append(f"{path.name}: Final Audit payload is not an array")
    if len(audit_records) != 200:
        failures.append(f"Final Audit baseline count is {len(audit_records)}, expected 200")
    if sorted(item.get("sequence") for item in audit_records) != list(range(1, 201)):
        failures.append("Final Audit sequences are not contiguous 1..200")

    record_files = sorted(RECORDS_PATH.glob("*.json")) if RECORDS_PATH.exists() else []
    records: list[dict[str, Any]] = []
    seen_sequences: set[int] = set()
    repeated: dict[str, defaultdict[str, list[str]]] = {
        field: defaultdict(list) for field in ANTI_TEMPLATE_FIELDS
    }
    for path in record_files:
        record = load_json_strict(path, failures, warnings)
        if not isinstance(record, dict):
            continue
        records.append(record)
        missing = sorted(REQUIRED_FIELDS - set(record))
        if missing:
            failures.append(f"{path.name}: missing fields {missing}")
        key = record.get("source_key")
        variant = record.get("case_variant_id")
        sequence = record.get("review_sequence")
        if sequence in seen_sequences:
            failures.append(f"duplicate record review sequence: {sequence}")
        seen_sequences.add(sequence)
        manifest_case = by_key.get(key)
        if not manifest_case:
            failures.append(f"{path.name}: source key not in manifest")
        else:
            if manifest_case.get("review_status") != "PENDING_FULL_REVIEW":
                failures.append(f"{path.name}: overwrites a completed baseline case")
            if variant != manifest_case.get("case_variant_id"):
                failures.append(f"{path.name}: case variant mismatch")
            if record.get("library_sequence") != manifest_case.get("library_sequence"):
                failures.append(f"{path.name}: library sequence mismatch")
            if record.get("source_integrity", {}).get("original_row_sha256") != manifest_case.get("original_row_sha256"):
                failures.append(f"{path.name}: original row hash mismatch")
        if record.get("automation_classification") not in CLASSES:
            failures.append(f"{path.name}: invalid automation classification")
        if record.get("review_outcome") not in OUTCOMES:
            failures.append(f"{path.name}: invalid review outcome")
        if record.get("second_review_outcome") not in SECOND_REVIEW_OUTCOMES:
            failures.append(f"{path.name}: invalid second-review outcome")
        if record.get("execution_status") != "NOT EXECUTED":
            failures.append(f"{path.name}: review task must remain NOT EXECUTED")
        if record.get("tool_result") != "NOT RUN":
            failures.append(f"{path.name}: review task must keep tool result NOT RUN")
        if record.get("engineer_verdict") != "ENGINEER OWNED; PENDING":
            failures.append(f"{path.name}: engineer verdict boundary changed")
        quality = record.get("per_case_quality_gate", {})
        for flag in (
            "exact_xlsx_row_read", "deepseek_same_row_read", "gpt_v2_read",
            "history_searched", "reasoning_is_case_specific",
            "source_identity_verified", "no_hardware_pass_claimed",
        ):
            if quality.get(flag) is not True:
                failures.append(f"{path.name}: failed quality gate {flag}")
        for field in ANTI_TEMPLATE_FIELDS:
            normalized = json.dumps(record.get(field), sort_keys=True, ensure_ascii=False)
            repeated[field][normalized].append(str(key))

    records.sort(key=lambda item: item.get("review_sequence", -1))
    if [item.get("review_sequence") for item in records] != list(range(201, 201 + len(records))):
        failures.append("new review sequences are not contiguous from 201")

    pending_cases = [case for case in cases if case.get("review_status") == "PENDING_FULL_REVIEW"]
    for index, record in enumerate(records):
        if index >= len(pending_cases):
            failures.append("more Full Review records than pending manifest cases")
            break
        expected_case = pending_cases[index]
        if (record.get("source_key"), record.get("case_variant_id")) != (
            expected_case.get("source_key"), expected_case.get("case_variant_id")
        ):
            failures.append(
                f"review order mismatch at sequence {record.get('review_sequence')}: "
                f"expected {expected_case.get('source_key')}"
            )

    combined = audit_records + records
    combined_keys = [item.get("source_key") for item in combined]
    combined_variants = [item.get("case_variant_id") for item in combined]
    if len(set(combined_keys)) != len(combined_keys):
        failures.append("effective records do not have unique Source Keys")
    if len(set(combined_variants)) != len(combined_variants):
        failures.append("effective records do not have unique Variant IDs")

    def effective_class(item: dict[str, Any]) -> Any:
        return item.get("automation_classification", item.get("after", {}).get("automation_classification"))

    combined_classes = Counter(effective_class(item) for item in combined)
    invalid_classes = sorted(str(value) for value in combined_classes if value not in CLASSES)
    if invalid_classes:
        failures.append(f"effective records contain invalid classifications: {invalid_classes}")
    expected_stats = {name: combined_classes.get(name, 0) for name in (
        "MANUAL ONLY", "REQUIRES PACKAGE / USER CONFIRMATION",
        "FULLY AUTOMATABLE", "BLOCKED",
    )}
    expected_completed = len(combined)
    if progress.get("new_full_review_records_count") != len(records):
        failures.append("progress new-record count does not match active record files")
    if progress.get("completed_count") != expected_completed:
        failures.append("progress completed count does not match baseline plus active records")
    if progress.get("remaining_count") != 3112 - expected_completed:
        failures.append("progress remaining count is inconsistent")
    if progress.get("classification_statistics") != expected_stats:
        failures.append(f"classification statistics mismatch: expected {expected_stats}, found {progress.get('classification_statistics')}")
    if sum(expected_stats.values()) != expected_completed:
        failures.append("derived classification total does not match completed count")
    if progress.get("classification_statistics_status") != "VERIFIED":
        failures.append("classification_statistics_status is not VERIFIED")

    expected_blocked = combined_classes.get("BLOCKED", 0)
    if progress.get("blocked_count") != expected_blocked:
        failures.append(f"blocked_count mismatch: expected {expected_blocked}")
    expected_conflicts = sum(
        bool(item.get("source_conflict_analysis", {}).get("is_source_conflict"))
        if "source_conflict_analysis" in item else bool(item.get("after", {}).get("source_conflict"))
        for item in combined
    )
    if progress.get("source_conflict_count") != expected_conflicts:
        failures.append(f"source_conflict_count mismatch: expected {expected_conflicts}")

    next_case = pending_cases[len(records)] if len(records) < len(pending_cases) else None
    expected_next_key = next_case.get("source_key") if next_case else None
    expected_next_variant = next_case.get("case_variant_id") if next_case else None
    if progress.get("next_unreviewed_source_key") != expected_next_key:
        failures.append(f"next Source Key mismatch: expected {expected_next_key}")
    if progress.get("next_unreviewed_case_variant_id") != expected_next_variant:
        failures.append(f"next Variant ID mismatch: expected {expected_next_variant}")

    report_text = RETROSPECTIVE_PATH.read_text(encoding="utf-8")
    report_rows = re.findall(r"^\| `([^`]+)` \| (ED-\d+) \| `([^`]+)` \|", report_text, flags=re.MULTILINE)
    if len(report_rows) != 36:
        failures.append(f"retrospective report has {len(report_rows)} affected rows, expected 36")
    record_by_key = {item.get("source_key"): item for item in records}
    overlay_keys = {item.get("source_key") for item in records if "engineering_decision_update" in item}
    report_keys = {key for key, _, _ in report_rows}
    if overlay_keys != report_keys:
        failures.append("engineering-decision overlay keys differ from retrospective report")
    for key, decision, classification in report_rows:
        record = record_by_key.get(key)
        if not record:
            failures.append(f"retrospective record missing: {key}")
            continue
        if record.get("automation_classification") != classification:
            failures.append(f"{key}: classification differs from retrospective report")
        overlay = record.get("engineering_decision_update", {})
        if decision not in overlay.get("decision_ids", []):
            failures.append(f"{key}: missing {decision} in engineering-decision overlay")
        status = overlay.get("engineering_decision_status", "")
        if decision == "ED-51":
            if "PENDING" not in status:
                failures.append(f"{key}: ED-51 must remain ENGINEER DECISION PENDING")
        elif status != "APPROVED AND APPLIED":
            failures.append(f"{key}: approved decision status is not applied")
        if overlay.get("review_count_change") != 0:
            failures.append(f"{key}: retrospective overlay changes completed count")

    decision_text = DECISIONS_PATH.read_text(encoding="utf-8")
    for number in range(33, 53):
        if f"ED-{number}" not in decision_text:
            failures.append(f"approved decision register missing ED-{number}")
    retrospective = progress.get("retrospective_audit_20261015", {})
    if retrospective.get("affected_record_count") != 36 or retrospective.get("completed_count_change") != 0:
        failures.append("progress retrospective audit count/change differs from approved report")
    if retrospective.get("remaining_engineer_decision_pending") != ["Functionality|708|Wistron-BIOS-00376-V002"]:
        failures.append("ED-51 pending identity differs from approved ruling")

    pending_decisions = progress.get("engineering_decisions_pending", [])
    pending_keys = [item.get("source_key") for item in pending_decisions if isinstance(item, dict)]
    if len(pending_keys) != len(set(pending_keys)):
        failures.append("engineering_decisions_pending contains duplicate Source Keys")
    unknown_pending = sorted(key for key in pending_keys if key not in set(combined_keys))
    if unknown_pending:
        failures.append(f"engineering_decisions_pending references non-effective records: {unknown_pending}")

    if progress.get("last_remote_checkpoint") != 891:
        failures.append("historical 891 checkpoint marker changed")
    checkpoint_exists = subprocess.run(
        ["git", "cat-file", "-e", f"{PUBLISHED_891_CHECKPOINT}^{{commit}}"], cwd=ROOT,
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False,
    ).returncode == 0
    if not checkpoint_exists:
        failures.append("published 891 checkpoint commit is unavailable")

    anti_template_groups = []
    for field, groups in repeated.items():
        for keys in groups.values():
            if len(keys) > 1:
                anti_template_groups.append({"field": field, "source_keys": keys})

    result = {
        "status": "PASS" if not failures else "FAIL",
        "manifest_count": len(cases),
        "completed_baseline_count": len(audit_records),
        "new_record_count": len(records),
        "effective_completed_count": expected_completed,
        "remaining_count": 3112 - expected_completed,
        "classification_statistics": expected_stats,
        "blocked_count": expected_blocked,
        "source_conflict_count": expected_conflicts,
        "engineering_decision_pending_count": len(pending_keys),
        "next_unreviewed_source_key": expected_next_key,
        "next_unreviewed_case_variant_id": expected_next_variant,
        "ed_33_52_affected_record_count": len(report_rows),
        "legacy_identical_duplicate_key_warning_count": sum(message.startswith("legacy identical duplicate key retained:") for message in warnings),
        "anti_template_exact_duplicate_group_count": len(anti_template_groups),
        "anti_template_exact_duplicate_candidates": anti_template_groups[:20],
        "warnings": warnings,
        "failures": failures,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not failures else 1


if __name__ == "__main__":
    raise SystemExit(main())

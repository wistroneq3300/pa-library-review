#!/usr/bin/env python3
"""Build the source-order 3,112-row identity manifest and initial progress.

This tool performs mechanical source/identity work only. It does not infer test
intent, assign a classification, or generate review conclusions.
"""

from __future__ import annotations

import hashlib
import json
import subprocess
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl


ROOT = Path(__file__).resolve().parents[3]
WORKSPACE = ROOT / "gpt-review" / "v3-1-full-3112"
XLSX = ROOT / "data" / "REVISED_commands_merged_with_raw.xlsx"
DEEPSEEK = ROOT / "data" / "tests.json"
GPT_V2 = ROOT / "gpt-review" / "test_gpt_review.json"
PILOT_MANIFEST = ROOT / "gpt-review" / "v3-pilot-200" / "SELECTION_MANIFEST.json"
FINAL_AUDIT_RECORDS = ROOT / "gpt-review" / "v3-1-final-audit" / "AUDIT_RECORDS"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def row_hash(headers: list[Any], values: tuple[Any, ...]) -> str:
    payload = dict(zip(headers, values, strict=True))
    serialized = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def load_final_audit() -> dict[str, dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for path in sorted(FINAL_AUDIT_RECORDS.glob("*.json")):
        records.extend(json.loads(path.read_text(encoding="utf-8")))
    if len(records) != 200:
        raise ValueError(f"Expected 200 Final Audit records, found {len(records)}")
    by_key = {record["source_key"]: record for record in records}
    if len(by_key) != 200:
        raise ValueError("Final Audit source keys are not unique")
    return by_key


def git_head() -> str:
    return subprocess.check_output(
        ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
    ).strip()


def main() -> None:
    pilot = json.loads(PILOT_MANIFEST.read_text(encoding="utf-8"))["cases"]
    pilot_by_key = {case["source_key"]: case for case in pilot}
    if len(pilot_by_key) != 200:
        raise ValueError("Pilot manifest does not contain 200 unique source keys")
    final_audit = load_final_audit()
    if set(final_audit) != set(pilot_by_key):
        raise ValueError("Final Audit identities do not exactly match the Pilot manifest")

    workbook = openpyxl.load_workbook(XLSX, read_only=True, data_only=False)
    cases: list[dict[str, Any]] = []
    by_sheet: Counter[str] = Counter()
    for sheet_name in workbook.sheetnames:
        rows = workbook[sheet_name].iter_rows(values_only=True)
        headers = list(next(rows))
        for physical_row, values in enumerate(rows, start=2):
            code = str(values[0] or "").strip()
            source_key = f"{sheet_name}|{physical_row}|{code}"
            case_variant_id = hashlib.sha256(source_key.encode("utf-8")).hexdigest()[:16]
            baseline = pilot_by_key.get(source_key)
            item = {
                "library_sequence": len(cases) + 1,
                "source_key": source_key,
                "case_variant_id": case_variant_id,
                "sheet": sheet_name,
                "source_row": physical_row,
                "code": code,
                "test_name": values[3],
                "sub_function": values[1],
                "test_set": values[2],
                "original_row_sha256": row_hash(headers, values),
                "review_status": "COMPLETED_FINAL_AUDIT_BASELINE" if baseline else "PENDING_FULL_REVIEW",
            }
            if baseline:
                audit = final_audit[source_key]
                if baseline["case_variant_id"] != case_variant_id:
                    raise ValueError(f"Pilot variant mismatch: {source_key}")
                if audit["case_variant_id"] != case_variant_id:
                    raise ValueError(f"Final Audit variant mismatch: {source_key}")
                item["final_audit_sequence"] = audit["sequence"]
                item["baseline_effective_classification"] = audit["after"]["automation_classification"]
                item["baseline_audit_outcome"] = audit["audit_outcome"]
            cases.append(item)
            by_sheet[sheet_name] += 1

    if len(cases) != 3112:
        raise ValueError(f"Expected 3,112 source rows, found {len(cases)}")
    if len({case["source_key"] for case in cases}) != 3112:
        raise ValueError("Full manifest source keys are not unique")
    if len({case["case_variant_id"] for case in cases}) != 3112:
        raise ValueError("Full manifest case_variant_id values are not unique")

    completed = [case for case in cases if case["review_status"].startswith("COMPLETED")]
    pending = [case for case in cases if case["review_status"].startswith("PENDING")]
    if len(completed) != 200 or len(pending) != 2912:
        raise ValueError(f"Expected 200/2,912 split, found {len(completed)}/{len(pending)}")

    expected_sheet_counts = {
        "Functionality": 2291,
        "Reliability": 202,
        "Performance": 60,
        "Compatibility": 459,
        "Stability": 88,
        "(No Main Function)": 12,
    }
    if dict(by_sheet) != expected_sheet_counts:
        raise ValueError(f"Unexpected sheet distribution: {dict(by_sheet)}")

    class_counts = Counter(case["baseline_effective_classification"] for case in completed)
    outcome_counts = Counter(case["baseline_audit_outcome"] for case in completed)
    manifest = {
        "schema_version": "v3-1-full-3112-selection-manifest-1.0",
        "generated_at": "2026-10-14T00:00:00+08:00",
        "branch": "codex/testcase-review-v3-1-full-3112",
        "baseline_final_audit_commit": "f337ec8e0a506e06cbd940d1781e7f484493d43c",
        "seq_182_decision_commit": git_head(),
        "source_sha256": {
            "xlsx": sha256(XLSX),
            "deepseek_tests_json": sha256(DEEPSEEK),
            "gpt_v2": sha256(GPT_V2),
            "pilot_selection_manifest": sha256(PILOT_MANIFEST),
        },
        "identity_method": "case_variant_id = first 16 lowercase hex characters of SHA-256(sheet|physical source_row|code)",
        "ordering": "Workbook sheet order, then physical XLSX source row order; the 200 completed baseline cases retain Final Audit sequence separately.",
        "total_count": 3112,
        "completed_baseline_count": 200,
        "pending_full_review_count": 2912,
        "sheet_distribution": dict(by_sheet),
        "baseline_effective_classification": dict(class_counts),
        "baseline_audit_outcomes": dict(outcome_counts),
        "cases": cases,
    }
    WORKSPACE.mkdir(parents=True, exist_ok=True)
    (WORKSPACE / "SELECTION_MANIFEST.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    progress = {
        "schema_version": "v3-1-full-3112-progress-1.0",
        "branch": "codex/testcase-review-v3-1-full-3112",
        "logic_version": "V3.1 Final Audit + ED-32 locked baseline",
        "target_count": 3112,
        "completed_count": 200,
        "new_full_review_records_count": 0,
        "remaining_count": 2912,
        "completed_baseline_reference": "gpt-review/v3-1-final-audit/AUDIT_RECORDS/",
        "next_unreviewed_source_key": pending[0]["source_key"],
        "next_unreviewed_case_variant_id": pending[0]["case_variant_id"],
        "classification_statistics": dict(class_counts),
        "blocked_count": class_counts.get("BLOCKED", 0),
        "source_conflict_count": 48,
        "engineering_decisions_pending": [],
        "project_readiness_inputs_pending": "Tracked per testcase; not counted as unresolved global engineering decisions.",
        "checkpoint_push_targets": [700, 1200, 1700, 2200, 2700, 3112],
        "last_remote_checkpoint": 200,
        "hardware_tested": False,
        "formal_test_verdicts_set": False,
    }
    (WORKSPACE / "PROGRESS.json").write_text(
        json.dumps(progress, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(
        json.dumps(
            {
                "manifest": str((WORKSPACE / "SELECTION_MANIFEST.json").relative_to(ROOT)),
                "progress": str((WORKSPACE / "PROGRESS.json").relative_to(ROOT)),
                "total": len(cases),
                "completed_baseline": len(completed),
                "pending": len(pending),
                "sheet_distribution": dict(by_sheet),
                "baseline_classification": dict(class_counts),
                "baseline_outcomes": dict(outcome_counts),
                "next_unreviewed_source_key": pending[0]["source_key"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
